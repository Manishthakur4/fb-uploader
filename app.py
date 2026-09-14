import os
import re
import json
import uuid
import csv
import hashlib
import time as _time
import threading
import requests
from datetime import datetime, timedelta
from flask import Flask, render_template, request, jsonify, redirect, url_for, session, send_from_directory
from functools import wraps
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from dotenv import load_dotenv
from werkzeug.utils import secure_filename

load_dotenv()

app = Flask(__name__)
app.secret_key = os.getenv("SECRET_KEY", "fb-uploader-secret-key-manish-2026")
app.config["UPLOAD_FOLDER"] = "uploads"
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024

ALLOWED_EXTENSIONS  = {"xlsx", "xls", "csv", "numbers"}
FB_GRAPH_URL        = "https://graph.facebook.com/v25.0"
PAGES_FILE          = "pages.json"
IMAGE_QUEUE_FILE    = "image_queue.json"
AUTH_EMAIL          = "manish@yopmail.com"
AUTH_PASSWORD       = "Manish@123"

TEXT_HEADERS   = {"text", "message", "post", "content", "post text", "posttext", "caption"}
UPLOAD_HEADERS = {"upload", "uploaded", "status", "done", "posted"}
DONE_VALUES    = {"yes", "uploaded", "done", "✓", "posted", "ok", "true", "yes ✓"}


# ── auth ──────────────────────────────────────────────────────────────────────

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated


# ── pages storage ─────────────────────────────────────────────────────────────

def load_pages():
    if not os.path.exists(PAGES_FILE):
        return []
    with open(PAGES_FILE) as f:
        return json.load(f)

def save_pages(pages):
    with open(PAGES_FILE, "w") as f:
        json.dump(pages, f, indent=2)


# ── file helpers ──────────────────────────────────────────────────────────────

def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS

def find_col(header_row, targets):
    for i, val in enumerate(header_row):
        if val and str(val).strip().lower() in targets:
            return i
    return None

def is_done(val):
    return str(val).strip().lower() in DONE_VALUES if val else False

def read_rows_from_file(filepath):
    ext = filepath.rsplit(".", 1)[-1].lower()

    if ext in ("xlsx", "xls"):
        wb       = load_workbook(filepath)
        ws       = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        if not all_rows:
            return [], None, 0
        header     = all_rows[0]
        text_col   = find_col(header, TEXT_HEADERS)
        if text_col is None:
            text_col = 1 if len(header) > 1 else 0
        upload_col = find_col(header, UPLOAD_HEADERS)
        rows = []
        for i, row in enumerate(all_rows[1:], start=2):
            val = str(row[text_col]).strip() if len(row) > text_col and row[text_col] is not None else ""
            if not val:
                continue
            up_val = row[upload_col] if upload_col is not None and len(row) > upload_col else None
            rows.append({"row_num": i, "text": val, "uploaded": is_done(up_val)})
        return rows, upload_col, text_col

    if ext == "csv":
        with open(filepath, newline="", encoding="utf-8-sig") as f:
            all_rows = list(csv.reader(f))
        if not all_rows:
            return [], None, 0
        header     = all_rows[0]
        text_col   = find_col(header, TEXT_HEADERS)
        if text_col is None:
            text_col = 1 if len(header) > 1 else 0
        upload_col = find_col(header, UPLOAD_HEADERS)
        rows = []
        for i, row in enumerate(all_rows[1:], start=2):
            val = row[text_col].strip() if len(row) > text_col else ""
            if not val:
                continue
            up_val = row[upload_col] if upload_col is not None and len(row) > upload_col else None
            rows.append({"row_num": i, "text": val, "uploaded": is_done(up_val)})
        return rows, upload_col, text_col

    if ext == "numbers":
        from numbers_parser import Document
        doc      = Document(filepath)
        table    = doc.sheets[0].tables[0]
        all_rows = list(table.iter_rows())
        if not all_rows:
            return [], None, 0
        header     = [c.value for c in all_rows[0]]
        text_col   = find_col(header, TEXT_HEADERS)
        if text_col is None:
            text_col = 1 if len(header) > 1 else 0
        upload_col = find_col(header, UPLOAD_HEADERS)
        rows = []
        for i, row in enumerate(all_rows[1:], start=2):
            val = str(row[text_col].value).strip() if len(row) > text_col and row[text_col].value is not None else ""
            if not val:
                continue
            up_val = row[upload_col].value if upload_col is not None and len(row) > upload_col else None
            rows.append({"row_num": i, "text": val, "uploaded": is_done(up_val)})
        return rows, upload_col, text_col

    return [], None, 0


def mark_row_uploaded(filepath, row_num, upload_col):
    ext   = filepath.rsplit(".", 1)[-1].lower()
    stamp = datetime.now().strftime("%d %b %Y %H:%M")
    if ext in ("xlsx", "xls"):
        wb   = load_workbook(filepath)
        ws   = wb.active
        if upload_col is not None:
            cell       = ws.cell(row=row_num, column=upload_col + 1)
            cell.value = f"Yes ✓  {stamp}"
            cell.font  = Font(color="1E7E34", bold=True)
            cell.fill  = PatternFill("solid", fgColor="E6F4EA")
        wb.save(filepath)
    elif ext == "csv":
        with open(filepath, newline="", encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
        if upload_col is not None and row_num - 1 < len(rows):
            while len(rows[row_num - 1]) <= upload_col:
                rows[row_num - 1].append("")
            rows[row_num - 1][upload_col] = f"Yes ✓  {stamp}"
        with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f).writerows(rows)


def post_to_facebook(text, page_id, access_token):
    url = f"{FB_GRAPH_URL}/{page_id}/feed"
    response = requests.post(url, data={"message": text, "access_token": access_token})
    return response.json()


# ── image queue helpers ───────────────────────────────────────────────────────

def load_image_queue():
    if not os.path.exists(IMAGE_QUEUE_FILE):
        return []
    with open(IMAGE_QUEUE_FILE) as f:
        return json.load(f)

def save_image_queue(queue):
    with open(IMAGE_QUEUE_FILE, "w") as f:
        json.dump(queue, f, indent=2)

def caption_from_url(url):
    path     = url.split("?")[0].rstrip("/")
    filename = path.split("/")[-1]
    name     = filename.rsplit(".", 1)[0] if "." in filename else filename
    return re.sub(r"[-_]+", " ", name).strip()

def post_image_to_facebook(img_url, caption, page_id, access_token):
    response = requests.post(
        f"{FB_GRAPH_URL}/{page_id}/photos",
        data={"url": img_url, "caption": caption, "access_token": access_token}
    )
    return response.json()


# ── cloudinary helpers ────────────────────────────────────────────────────────

def _cld_sign(params):
    secret      = os.getenv("CLOUDINARY_API_SECRET", "")
    sorted_str  = "&".join(f"{k}={v}" for k, v in sorted(params.items()))
    return hashlib.sha1(f"{sorted_str}{secret}".encode()).hexdigest()

def cloudinary_upload(file_stream, original_filename):
    cloud   = os.getenv("CLOUDINARY_CLOUD_NAME", "")
    api_key = os.getenv("CLOUDINARY_API_KEY", "")
    ts      = int(_time.time())
    folder  = "fb-uploader"
    params  = {"folder": folder, "timestamp": ts}
    res = requests.post(
        f"https://api.cloudinary.com/v1_1/{cloud}/image/upload",
        data={"api_key": api_key, "timestamp": ts,
              "signature": _cld_sign(params), "folder": folder},
        files={"file": (original_filename, file_stream)},
        timeout=30,
    )
    return res.json()

def cloudinary_delete(public_id):
    cloud   = os.getenv("CLOUDINARY_CLOUD_NAME", "")
    api_key = os.getenv("CLOUDINARY_API_KEY", "")
    ts      = int(_time.time())
    params  = {"public_id": public_id, "timestamp": ts}
    res = requests.post(
        f"https://api.cloudinary.com/v1_1/{cloud}/image/destroy",
        data={"api_key": api_key, "timestamp": ts,
              "signature": _cld_sign(params), "public_id": public_id},
        timeout=15,
    )
    return res.json()


# ── multi-job text scheduler ──────────────────────────────────────────────────

_jobs_lock = threading.Lock()
_jobs      = {}   # job_id → state dict
_timers    = {}   # job_id → threading.Timer


def _make_job(job_id, filename, page_ids, interval_min, total):
    pages      = load_pages()
    page_names = [p["name"] for p in pages if p["id"] in page_ids]
    return {
        "id":           job_id,
        "filename":     filename,
        "page_ids":     page_ids,
        "page_names":   page_names,
        "interval_min": interval_min,
        "status":       "running",
        "posts_sent":   0,
        "posts_failed": 0,
        "total":        total,
        "next_at":      None,
        "last_post":    None,
        "uploaded_rows":[],
        "created_at":   datetime.now().strftime("%d %b %Y %H:%M"),
    }


def _check_queue(finished_page_ids):
    """After a job finishes/stops, start the next queued job waiting on any of those pages."""
    next_job_id = None
    with _jobs_lock:
        for jid, job in _jobs.items():
            if job["status"] == "queued" and any(pid in job["page_ids"] for pid in finished_page_ids):
                job["status"] = "running"
                next_job_id   = jid
                break
    if next_job_id:
        t = threading.Thread(target=_tick, args=[next_job_id], daemon=True)
        t.start()


def _tick(job_id):
    with _jobs_lock:
        if job_id not in _jobs or _jobs[job_id]["status"] != "running":
            return
        job          = _jobs[job_id]
        filename     = job["filename"]
        page_ids     = job["page_ids"]
        interval_min = job["interval_min"]

    filepath = os.path.join(app.config["UPLOAD_FOLDER"], secure_filename(filename))

    try:
        rows, upload_col, _ = read_rows_from_file(filepath)
    except Exception as e:
        with _jobs_lock:
            if job_id in _jobs:
                _jobs[job_id]["last_post"] = {"status": "error", "error": str(e), "time": datetime.now().strftime("%d %b %Y %H:%M:%S")}
        return

    pending = [r for r in rows if not r["uploaded"]]
    if not pending:
        with _jobs_lock:
            if job_id in _jobs:
                _jobs[job_id]["status"]  = "done"
                _jobs[job_id]["next_at"] = None
                finished_page_ids        = _jobs[job_id]["page_ids"][:]
        _check_queue(finished_page_ids)
        return

    row       = pending[0]
    pages     = load_pages()
    selected  = [p for p in pages if p["id"] in page_ids]
    success   = False
    post_id   = None
    page_name = selected[0]["name"] if selected else ""
    error_msg = ""

    for page in selected:
        result    = post_to_facebook(row["text"], page["page_id"], page["access_token"])
        page_name = page["name"]
        if "id" in result:
            success = True
            post_id = result["id"]
        else:
            error_msg = result.get("error", {}).get("message", "Unknown error")

    now = datetime.now()

    if success:
        try:
            mark_row_uploaded(filepath, row["row_num"], upload_col)
        except Exception:
            pass
        with _jobs_lock:
            if job_id in _jobs:
                _jobs[job_id]["posts_sent"]    += 1
                _jobs[job_id]["uploaded_rows"].append(row["row_num"])
                _jobs[job_id]["last_post"] = {
                    "row_num": row["row_num"], "text": row["text"][:80],
                    "page": page_name, "status": "success",
                    "post_id": post_id, "time": now.strftime("%d %b %Y %H:%M:%S"),
                }
    else:
        with _jobs_lock:
            if job_id in _jobs:
                _jobs[job_id]["posts_failed"] += 1
                _jobs[job_id]["last_post"] = {
                    "row_num": row["row_num"], "text": row["text"][:80],
                    "page": page_name, "status": "failed",
                    "error": error_msg, "time": now.strftime("%d %b %Y %H:%M:%S"),
                }

    try:
        rows2, _, _ = read_rows_from_file(filepath)
        still_pending = [r for r in rows2 if not r["uploaded"]]
    except Exception:
        still_pending = []

    finished_page_ids = []
    with _jobs_lock:
        if job_id not in _jobs:
            return
        if not still_pending:
            _jobs[job_id]["status"]  = "done"
            _jobs[job_id]["next_at"] = None
            finished_page_ids        = _jobs[job_id]["page_ids"][:]
        elif _jobs[job_id]["status"] == "running":
            next_dt = now + timedelta(minutes=interval_min)
            _jobs[job_id]["next_at"] = next_dt.isoformat()
            t = threading.Timer(interval_min * 60, _tick, args=[job_id])
            t.daemon = True
            _timers[job_id] = t
            t.start()

    if finished_page_ids:
        _check_queue(finished_page_ids)


# ── routes ─────────────────────────────────────────────────────────────────────

@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("logged_in"):
        return redirect(url_for("index"))
    error = None
    if request.method == "POST":
        email    = request.form.get("email", "").strip()
        password = request.form.get("password", "")
        if email == AUTH_EMAIL and password == AUTH_PASSWORD:
            session["logged_in"] = True
            return redirect(url_for("index"))
        error = "Invalid email or password."
    return render_template("login.html", error=error)

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

@app.route("/")
@login_required
def index():
    return render_template("index.html")


# ── token exchange ────────────────────────────────────────────────────────────

@app.route("/token/exchange", methods=["POST"])
@login_required
def token_exchange():
    data       = request.get_json()
    user_token = data.get("user_token", "").strip()
    app_id     = data.get("app_id", "").strip() or os.getenv("FB_APP_ID", "")
    app_secret = data.get("app_secret", "").strip() or os.getenv("FB_APP_SECRET", "")

    if not user_token:
        return jsonify({"error": "Paste your User Access Token from Graph API Explorer."}), 400
    if not app_id or not app_secret:
        return jsonify({"error": "App ID and App Secret are required. Get them from developers.facebook.com → your app → Settings → Basic."}), 400

    # Step 1: exchange short-lived user token → long-lived user token (60 days)
    ex = requests.get(f"{FB_GRAPH_URL}/oauth/access_token", params={
        "grant_type":        "fb_exchange_token",
        "client_id":         app_id,
        "client_secret":     app_secret,
        "fb_exchange_token": user_token,
    })
    ex_data = ex.json()
    if "error" in ex_data:
        return jsonify({"error": "Exchange failed: " + ex_data["error"].get("message", "Unknown error")}), 400

    long_token = ex_data.get("access_token")

    # Step 2: fetch all pages the user manages with their permanent page tokens
    accts = requests.get(f"{FB_GRAPH_URL}/me/accounts", params={
        "access_token": long_token,
        "fields":       "id,name,access_token",
    })
    accts_data = accts.json()
    if "error" in accts_data:
        return jsonify({"error": "Could not fetch pages: " + accts_data["error"].get("message", "")}), 400

    fb_pages = accts_data.get("data", [])
    if not fb_pages:
        return jsonify({"error": "No pages found for this account. Make sure you manage at least one Facebook Page."}), 400

    # Step 3: update existing pages and add new ones with never-expiring tokens
    saved   = load_pages()
    updated = 0
    added   = []
    for fb_p in fb_pages:
        matched = False
        for sp in saved:
            if sp["page_id"] == fb_p["id"]:
                sp["access_token"] = fb_p["access_token"]
                sp["name"]         = fb_p["name"]
                matched = True
                updated += 1
                break
        if not matched:
            saved.append({
                "id":           str(uuid.uuid4()),
                "name":         fb_p["name"],
                "page_id":      fb_p["id"],
                "access_token": fb_p["access_token"],
            })
            added.append(fb_p["name"])

    save_pages(saved)

    return jsonify({
        "success":       True,
        "pages_found":   [{"name": p["name"], "id": p["id"]} for p in fb_pages],
        "pages_updated": updated,
        "pages_added":   len(added),
        "added_names":   added,
    })


# ── pages API ─────────────────────────────────────────────────────────────────

@app.route("/pages", methods=["GET"])
@login_required
def get_pages():
    pages = load_pages()
    return jsonify([{"id": p["id"], "name": p["name"], "page_id": p["page_id"]} for p in pages])

@app.route("/pages/add", methods=["POST"])
@login_required
def add_page():
    data         = request.get_json()
    name         = data.get("name", "").strip()
    page_id      = data.get("page_id", "").strip()
    access_token = data.get("access_token", "").strip()
    if not name or not page_id or not access_token:
        return jsonify({"error": "Name, Page ID, and Access Token are all required."}), 400
    pages = load_pages()
    pages.append({"id": str(uuid.uuid4()), "name": name, "page_id": page_id, "access_token": access_token})
    save_pages(pages)
    return jsonify({"success": True})

@app.route("/pages/delete/<pid>", methods=["DELETE"])
@login_required
def delete_page(pid):
    save_pages([p for p in load_pages() if p["id"] != pid])
    return jsonify({"success": True})


# ── upload ─────────────────────────────────────────────────────────────────────

@app.route("/preview", methods=["POST"])
@login_required
def preview():
    if "file" not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "No file selected"}), 400
    if not allowed_file(file.filename):
        return jsonify({"error": "Unsupported file type. Use xlsx, xls, csv, or numbers."}), 400

    filename = secure_filename(file.filename)
    filepath = os.path.join(app.config["UPLOAD_FOLDER"], filename)
    file.save(filepath)

    try:
        rows, upload_col, _ = read_rows_from_file(filepath)
    except Exception as e:
        return jsonify({"error": f"Could not read file: {str(e)}"}), 400

    if not rows:
        return jsonify({"error": "No post text found. Make sure your file has a 'Post Text' column."}), 400

    pending  = [r for r in rows if not r["uploaded"]]
    uploaded = [r for r in rows if r["uploaded"]]
    return jsonify({
        "filename":      filename,
        "rows":          rows,
        "total":         len(rows),
        "pending_count": len(pending),
        "done_count":    len(uploaded),
    })


# ── scheduler API ─────────────────────────────────────────────────────────────

@app.route("/scheduler/start", methods=["POST"])
@login_required
def scheduler_start():
    data         = request.get_json()
    filename     = data.get("filename", "")
    page_ids     = data.get("page_ids", [])
    interval_min = int(data.get("interval_min", 20))

    if not filename:
        return jsonify({"error": "No file selected."}), 400
    if not page_ids:
        return jsonify({"error": "Select at least one Facebook page."}), 400

    filepath = os.path.join(app.config["UPLOAD_FOLDER"], secure_filename(filename))
    try:
        rows, _, _ = read_rows_from_file(filepath)
    except Exception as e:
        return jsonify({"error": str(e)}), 400

    pending = [r for r in rows if not r["uploaded"]]
    if not pending:
        return jsonify({"error": "No pending posts — all rows are already uploaded."}), 400

    # Check if any selected page already has a running job → queue this one
    with _jobs_lock:
        busy_pages = set()
        for job in _jobs.values():
            if job["status"] == "running":
                busy_pages.update(job["page_ids"])
        is_busy = any(pid in busy_pages for pid in page_ids)

    job_id = str(uuid.uuid4())
    with _jobs_lock:
        job = _make_job(job_id, filename, page_ids, interval_min, len(pending))
        if is_busy:
            job["status"] = "queued"
        _jobs[job_id] = job

    if not is_busy:
        t = threading.Thread(target=_tick, args=[job_id], daemon=True)
        t.start()

    return jsonify({"success": True, "job_id": job_id, "queued": is_busy})


@app.route("/scheduler/stop/<job_id>", methods=["POST"])
@login_required
def scheduler_stop(job_id):
    finished_page_ids = []
    with _jobs_lock:
        if job_id in _timers:
            _timers[job_id].cancel()
            del _timers[job_id]
        if job_id in _jobs:
            _jobs[job_id]["status"]  = "stopped"
            _jobs[job_id]["next_at"] = None
            finished_page_ids        = _jobs[job_id]["page_ids"][:]
    _check_queue(finished_page_ids)
    return jsonify({"success": True})


@app.route("/scheduler/remove/<job_id>", methods=["DELETE"])
@login_required
def scheduler_remove(job_id):
    with _jobs_lock:
        if job_id in _timers:
            _timers[job_id].cancel()
            del _timers[job_id]
        _jobs.pop(job_id, None)
    return jsonify({"success": True})


@app.route("/scheduler/status", methods=["GET"])
@login_required
def scheduler_status():
    with _jobs_lock:
        return jsonify(list(_jobs.values()))


# ── download ──────────────────────────────────────────────────────────────────

@app.route("/download/<filename>")
@login_required
def download_file(filename):
    return send_from_directory(app.config["UPLOAD_FOLDER"], secure_filename(filename), as_attachment=True)


# ── image queue routes ────────────────────────────────────────────────────────

@app.route("/images", methods=["GET"])
@login_required
def get_images():
    return jsonify(load_image_queue())

@app.route("/images/add", methods=["POST"])
@login_required
def add_images():
    data = request.get_json()
    urls = data.get("urls", [])
    if not urls:
        return jsonify({"error": "No URLs provided"}), 400
    queue   = load_image_queue()
    existing = {img["url"] for img in queue}
    added   = 0
    for url in urls:
        url = url.strip()
        if not url or url in existing:
            continue
        queue.append({
            "id":        str(uuid.uuid4()),
            "url":       url,
            "public_id": None,
            "caption":   caption_from_url(url),
            "status":    "pending",
            "posted_at": None,
        })
        existing.add(url)
        added += 1
    save_image_queue(queue)
    return jsonify({"success": True, "added": added})

@app.route("/images/<img_id>/delete", methods=["DELETE"])
@login_required
def delete_image(img_id):
    save_image_queue([img for img in load_image_queue() if img["id"] != img_id])
    return jsonify({"success": True})

@app.route("/images/<img_id>/caption", methods=["PATCH"])
@login_required
def update_image_caption(img_id):
    caption = request.get_json().get("caption", "").strip()
    queue   = load_image_queue()
    for img in queue:
        if img["id"] == img_id:
            img["caption"] = caption
            break
    save_image_queue(queue)
    return jsonify({"success": True})

@app.route("/images/upload", methods=["POST"])
@login_required
def upload_image_to_cloudinary():
    if "file" not in request.files:
        return jsonify({"error": "No file provided"}), 400
    file = request.files["file"]
    if not file.filename:
        return jsonify({"error": "Empty filename"}), 400
    ext = file.filename.rsplit(".", 1)[-1].lower()
    if ext not in {"jpg", "jpeg", "png", "gif", "webp", "heif", "tiff"}:
        return jsonify({"error": f"Unsupported format: .{ext}. Use JPG, PNG, GIF or WebP."}), 400

    result = cloudinary_upload(file.stream, secure_filename(file.filename))
    if "error" in result:
        return jsonify({"error": result["error"].get("message", "Cloudinary upload failed")}), 400

    url        = result["secure_url"]
    public_id  = result["public_id"]
    # Use the original filename (without extension) as the caption
    base_name  = file.filename.rsplit(".", 1)[0]
    caption    = re.sub(r"[-_]+", " ", base_name).strip()

    queue = load_image_queue()
    if not any(img.get("public_id") == public_id for img in queue):
        queue.append({
            "id":        str(uuid.uuid4()),
            "url":       url,
            "public_id": public_id,
            "caption":   caption,
            "status":    "pending",
            "posted_at": None,
        })
        save_image_queue(queue)

    return jsonify({"success": True, "url": url, "public_id": public_id, "caption": caption})


# ── image scheduler ───────────────────────────────────────────────────────────

_img_lock   = threading.Lock()
_img_jobs   = {}
_img_timers = {}


def _make_img_job(job_id, page_ids, images_per_day):
    pages        = load_pages()
    page_names   = [p["name"] for p in pages if p["id"] in page_ids]
    interval_min = max(1, int((24 * 60) / images_per_day))
    return {
        "id":             job_id,
        "page_ids":       page_ids,
        "page_names":     page_names,
        "images_per_day": images_per_day,
        "interval_min":   interval_min,
        "status":         "running",
        "posts_sent":     0,
        "posts_failed":   0,
        "next_at":        None,
        "last_post":      None,
        "created_at":     datetime.now().strftime("%d %b %Y %H:%M"),
    }


def _check_img_queue(finished_page_ids):
    next_job_id = None
    with _img_lock:
        for jid, job in _img_jobs.items():
            if job["status"] == "queued" and any(pid in job["page_ids"] for pid in finished_page_ids):
                job["status"] = "running"
                next_job_id   = jid
                break
    if next_job_id:
        t = threading.Thread(target=_img_tick, args=[next_job_id], daemon=True)
        t.start()


def _img_tick(job_id):
    with _img_lock:
        if job_id not in _img_jobs or _img_jobs[job_id]["status"] != "running":
            return
        job          = _img_jobs[job_id]
        page_ids     = job["page_ids"]
        interval_min = job["interval_min"]

    queue   = load_image_queue()
    pending = [img for img in queue if img["status"] == "pending"]

    if not pending:
        finished_page_ids = []
        with _img_lock:
            if job_id in _img_jobs:
                _img_jobs[job_id]["status"]  = "done"
                _img_jobs[job_id]["next_at"] = None
                finished_page_ids            = _img_jobs[job_id]["page_ids"][:]
        _check_img_queue(finished_page_ids)
        return

    img      = pending[0]
    pages    = load_pages()
    selected = [p for p in pages if p["id"] in page_ids]
    success   = False
    post_id   = None
    page_name = selected[0]["name"] if selected else ""
    error_msg = ""

    for page in selected:
        result    = post_image_to_facebook(img["url"], img["caption"], page["page_id"], page["access_token"])
        page_name = page["name"]
        if "id" in result:
            success = True
            post_id = result["id"]
        else:
            error_msg = result.get("error", {}).get("message", "Unknown error")

    now = datetime.now()

    if success:
        # Delete from Cloudinary now that it's posted to Facebook
        if img.get("public_id"):
            try:
                cloudinary_delete(img["public_id"])
            except Exception:
                pass

        # Remove from queue entirely (already posted + deleted from cloud)
        queue_updated = load_image_queue()
        save_image_queue([q for q in queue_updated if q["id"] != img["id"]])

        with _img_lock:
            if job_id in _img_jobs:
                _img_jobs[job_id]["posts_sent"] += 1
                _img_jobs[job_id]["last_post"]   = {
                    "url": img["url"], "caption": img["caption"],
                    "page": page_name, "status": "success",
                    "post_id": post_id, "time": now.strftime("%d %b %Y %H:%M:%S"),
                }
    else:
        with _img_lock:
            if job_id in _img_jobs:
                _img_jobs[job_id]["posts_failed"] += 1
                _img_jobs[job_id]["last_post"]     = {
                    "url": img["url"], "caption": img["caption"],
                    "page": page_name, "status": "failed",
                    "error": error_msg, "time": now.strftime("%d %b %Y %H:%M:%S"),
                }

    still_pending = [i for i in load_image_queue() if i["status"] == "pending"]
    finished_page_ids = []

    with _img_lock:
        if job_id not in _img_jobs:
            return
        if not still_pending:
            _img_jobs[job_id]["status"]  = "done"
            _img_jobs[job_id]["next_at"] = None
            finished_page_ids            = _img_jobs[job_id]["page_ids"][:]
        elif _img_jobs[job_id]["status"] == "running":
            next_dt = now + timedelta(minutes=interval_min)
            _img_jobs[job_id]["next_at"] = next_dt.isoformat()
            t = threading.Timer(interval_min * 60, _img_tick, args=[job_id])
            t.daemon = True
            _img_timers[job_id] = t
            t.start()

    if finished_page_ids:
        _check_img_queue(finished_page_ids)


@app.route("/images/scheduler/start", methods=["POST"])
@login_required
def img_scheduler_start():
    data           = request.get_json()
    page_ids       = data.get("page_ids", [])
    images_per_day = int(data.get("images_per_day", 2))

    if not page_ids:
        return jsonify({"error": "Select at least one Facebook page."}), 400

    pending = [img for img in load_image_queue() if img["status"] == "pending"]
    if not pending:
        return jsonify({"error": "No pending images in the queue. Add some URLs first."}), 400

    with _img_lock:
        busy_pages = {pid for job in _img_jobs.values() if job["status"] == "running" for pid in job["page_ids"]}
        is_busy    = any(pid in busy_pages for pid in page_ids)

    job_id = str(uuid.uuid4())
    with _img_lock:
        job = _make_img_job(job_id, page_ids, images_per_day)
        if is_busy:
            job["status"] = "queued"
        _img_jobs[job_id] = job

    if not is_busy:
        t = threading.Thread(target=_img_tick, args=[job_id], daemon=True)
        t.start()

    return jsonify({"success": True, "job_id": job_id, "queued": is_busy})


@app.route("/images/scheduler/stop/<job_id>", methods=["POST"])
@login_required
def img_scheduler_stop(job_id):
    finished_page_ids = []
    with _img_lock:
        if job_id in _img_timers:
            _img_timers[job_id].cancel()
            del _img_timers[job_id]
        if job_id in _img_jobs:
            _img_jobs[job_id]["status"]  = "stopped"
            _img_jobs[job_id]["next_at"] = None
            finished_page_ids            = _img_jobs[job_id]["page_ids"][:]
    _check_img_queue(finished_page_ids)
    return jsonify({"success": True})


@app.route("/images/scheduler/remove/<job_id>", methods=["DELETE"])
@login_required
def img_scheduler_remove(job_id):
    with _img_lock:
        if job_id in _img_timers:
            _img_timers[job_id].cancel()
            del _img_timers[job_id]
        _img_jobs.pop(job_id, None)
    return jsonify({"success": True})


@app.route("/images/scheduler/status", methods=["GET"])
@login_required
def img_scheduler_status():
    with _img_lock:
        return jsonify(list(_img_jobs.values()))


if __name__ == "__main__":
    os.makedirs("uploads", exist_ok=True)
    app.run(debug=True, port=8888)
